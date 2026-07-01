import csv
import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from typing import List, Dict

class ReportGenerator:
    @staticmethod
    def generate_csv(records: List[Dict]) -> str:
        """
        Generates standard CSV content.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Headers
        writer.writerow(["employee_id", "employee_name", "date", "check_in", "check_out", "status", "total_hours"])
        
        # Rows
        for row in records:
            writer.writerow([
                row.get("employee_id"),
                row.get("employee_name"),
                row.get("date"),
                row.get("check_in", ""),
                row.get("check_out", ""),
                row.get("status"),
                row.get("total_hours", 0.0)
            ])
            
        return output.getvalue()

    @staticmethod
    def generate_excel(records: List[Dict]) -> io.BytesIO:
        """
        Generates formatted Excel workbook.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"
        
        # Title Block
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "AI Face Recognition Attendance System Report"
        title_cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="005FAF", end_color="005FAF", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 40
        
        # Headers
        headers = ["Employee ID", "Employee Name", "Date", "Check-In", "Check-Out", "Status", "Hours Worked"]
        ws.append([]) # Spacer row
        ws.append(headers)
        
        # Format Headers
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="535F70", end_color="535F70", fill_type="solid")
        for col_num in range(1, 8):
            cell = ws.cell(row=3, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            
        # Append data
        for row in records:
            ws.append([
                row.get("employee_id"),
                row.get("employee_name"),
                row.get("date"),
                row.get("check_in", ""),
                row.get("check_out", ""),
                row.get("status"),
                row.get("total_hours", 0.0)
            ])
            
        # Apply conditional coloring for status
        green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        orange_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        
        for r_idx in range(4, len(records) + 4):
            status_cell = ws.cell(row=r_idx, column=6)
            status = str(status_cell.value)
            if status in ["PRESENT", "REGULAR"]:
                status_cell.fill = green_fill
            elif status in ["LATE", "EARLY_LEAVE"]:
                status_cell.fill = orange_fill
            elif status in ["ABSENT", "INCOMPLETE"]:
                status_cell.fill = red_fill

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream

    @staticmethod
    def generate_pdf(records: List[Dict]) -> io.BytesIO:
        """
        Generates formatted PDF report.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        story = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=20,
            textColor=colors.HexColor("#005FAF"),
            spaceAfter=20,
            alignment=1 # Center
        )
        
        story.append(Paragraph("AI Face Recognition Attendance System Report", title_style))
        story.append(Spacer(1, 10))
        
        # Table data assembly
        table_data = [["Employee ID", "Employee Name", "Date", "Check-In", "Check-Out", "Status"]]
        for row in records:
            table_data.append([
                row.get("employee_id"),
                row.get("employee_name"),
                row.get("date"),
                row.get("check_in", ""),
                row.get("check_out", ""),
                row.get("status")
            ])
            
        t = Table(table_data, colWidths=[80, 120, 80, 80, 80, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#535F70")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor("#F0F4F8")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#BBC7DB")),
            ('FONTSIZE', (0,0), (-1,-1), 9),
        ]))
        
        story.append(t)
        doc.build(story)
        buffer.seek(0)
        return buffer
